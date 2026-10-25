import docx
import re

doc = docx.Document('d:/trienkhaihethong/report/BAO_CAO_CUOI_KY.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Total Sections: {len(doc.sections)}")

total_words = 0
placeholders = []
page_breaks = 0
images_count = 0

for p in doc.paragraphs:
    text = p.text
    total_words += len(text.split())
    if "PLACEHOLDER" in text or re.search(r'\[.*?\]', text):
        placeholders.append(text.strip())
    for r in p.runs:
        if 'w:br' in r._r.xml and 'type="page"' in r._r.xml:
            page_breaks += 1
        if 'w:drawing' in r._r.xml:
            images_count += 1

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                total_words += len(p.text.split())
                if "PLACEHOLDER" in p.text:
                    placeholders.append(p.text.strip())

print(f"Total Words: {total_words}")
print(f"Explicit Page Breaks: {page_breaks}")
print(f"Images count: {images_count}")
print(f"\nPlaceholders found ({len(placeholders)} instances):")
for ph in list(set(placeholders)):
    print(f"  - {ph}")
